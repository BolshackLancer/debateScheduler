from flask import Flask, render_template, request, jsonify, Response
import json
import csv
import openpyxl
from typing import List, Tuple, BinaryIO
from scheduler import generate_unique_rounds
app = Flask(__name__)


def generateRounds(numRounds, numRooms, roster):
    # 1. Parse roster to get team lists
    cubs_teams, lions_teams = parse_roster(roster)
    # 2. Print cubs teams and lions teams separately
    print("Cubs Teams:", cubs_teams)
    print("Lions Teams:", lions_teams)
    # 3. Create cubs rounds
    cubs_rounds = generate_unique_rounds(cubs_teams, int(numRounds))
    lions_rounds = generate_unique_rounds(lions_teams, int(numRounds))
    # 4. Write cubs rounds to CSV
    import csv
    from io import StringIO
    cubs_csv = StringIO()
    lions_csv = StringIO()
    # Write Cubs rounds
    max_cubs_room = 0
    for round_num, round_matches in enumerate(cubs_rounds, 1):
        cubs_csv.write(f"Round {round_num}\n")
        for room_num, match in enumerate(round_matches, 1):
            t1, t2 = match
            if t2 is None or t2 == "BYE":
                cubs_csv.write(f"No Room,{t1},vs,BYE\n")
            else:
                actual_room = room_num
                cubs_csv.write(f"Room {actual_room},{t1},vs,{t2}\n")
                if actual_room > max_cubs_room:
                    max_cubs_room = actual_room
        cubs_csv.write("\n")
    # Write Lions rounds, room numbers resume from max_cubs_room + 1
    for round_num, round_matches in enumerate(lions_rounds, 1):
        lions_csv.write(f"Round {round_num}\n")
        for room_num, match in enumerate(round_matches, 1):
            t1, t2 = match
            if t2 is None or t2 == "BYE":
                lions_csv.write(f"No Room,{t1},vs,BYE\n")
            else:
                actual_room = max_cubs_room + room_num
                lions_csv.write(f"Room {actual_room},{t1},vs,{t2}\n")
        lions_csv.write("\n")
    # 7. Return combined CSV
    combined_csv = cubs_csv.getvalue() + "\n" + lions_csv.getvalue()
    print(f"Generated rounds for Cubs and Lions successfully: {combined_csv}")
    return combined_csv
	


@app.route('/')
def index():
	return render_template('upload.html')

@app.route('/createRounds', methods=['POST'])
def createRounds():
    print("entered createRounds")
    numRounds = request.form['numRounds']
    numRooms = request.form['numRooms']
    print(numRounds)
    roster = request.files['roster']
    filename = roster.filename
    print(filename)
    generated_csv = generateRounds(numRounds, numRooms, roster)
    response = Response(
        generated_csv,
        mimetype="text/csv",
        headers={"Content-Disposition": "attachment;filename=rounds.csv"}
    )
    return response

@app.route('/raw_rounds', methods=['POST'])
def raw_rounds():
    # Expecting a form field 'teams' with comma separated team names
    teams_str = request.form.get('teams', '')
    team_names = [t.strip() for t in teams_str.split(',') if t.strip()]
    rounds = generate_unique_rounds(team_names, num_rounds=3)
    # Format the rounds for JSON response
    rounds_json = []
    for rnd in rounds:
        round_list = []
        for match in rnd:
            t1, t2 = match
            round_list.append({'team1': t1, 'team2': t2})
        rounds_json.append(round_list)
    return jsonify({'rounds': rounds_json})

def parse_roster(roster_file: BinaryIO) -> Tuple[List[str], List[str]]:
    """
    Parse the uploaded Excel roster and return two lists:
    - cubs_teams: unique team names from 'Cubs Registration' with at least 2 debater names
    - lions_teams: unique team names from 'Lions Registration' with at least 2 debater names
    """
    wb = openpyxl.load_workbook(roster_file, data_only=True)
    cubs_teams = []
    lions_teams = []
    for sheet_name, team_list in [('Cubs Registration', cubs_teams), ('Lions Registration', lions_teams)]:
        if sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            row = 5
            while True:
                team_cell = ws[f'C{row}']
                team_name = team_cell.value
                if not team_name:
                    break
                debaters = [ws[f'F{row+i}'].value for i in range(3)]
                valid_debaters = [d for d in debaters if d and str(d).strip()]
                if len(valid_debaters) >= 2 and team_name not in team_list:
                    team_list.append(str(team_name).strip())
                row += 3
    return cubs_teams, lions_teams

if __name__ == '__main__':
	app.run()